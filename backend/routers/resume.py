from fastapi import APIRouter, UploadFile, File, Form, HTTPException, Depends
from pydantic import BaseModel, Field
from typing import List
import fitz # PyMuPDF
import docx
import io

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.prompts import PromptTemplate
from langchain_core.output_parsers import PydanticOutputParser
from core.config import settings
from core.auth import get_current_user, UserData
from core.db import get_auth_client

router = APIRouter()

class AnalyzeRequest(BaseModel):
    job_description: str

class ResumeAnalysis(BaseModel):
    company_name: str = Field(description="The name of the company hiring for the position, if mentioned. Otherwise 'Unknown Company'.")
    job_title: str = Field(description="The job title from the job description.")
    ats_score: int = Field(description="The ATS score out of 100 based on the match between resume and job description.")
    matching_skills: List[str] = Field(description="A list of matching skills found in both the resume and job description.")
    missing_skills: List[str] = Field(description="A list of important skills required in the job description that are missing in the resume.")
    summary: str = Field(description="A short summary detailing the overall fit and specific areas for improvement.")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = fitz.open("pdf", file_bytes)
        for page in doc:
            text += page.get_text()
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {str(e)}")
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")
    return text

@router.post("/analyze")
async def analyze_resume(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
    current_user: UserData = Depends(get_current_user)
):
    if not resume.filename.endswith(('.pdf', '.docx')):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported.")
    
    file_bytes = await resume.read()
    text = ""
    
    if resume.filename.endswith('.pdf'):
        try:
            text = extract_text_from_pdf(file_bytes)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
    elif resume.filename.endswith('.docx'):
        try:
            text = extract_text_from_docx(file_bytes)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))

    if not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from the provided document.")

    # Initialize Gemini Chat Model
    llm = ChatGoogleGenerativeAI(
        model="gemini-2.5-flash",
        temperature=0.2,
        google_api_key=settings.GEMINI_API_KEY
    )

    # Setup Output Parser
    parser = PydanticOutputParser(pydantic_object=ResumeAnalysis)

    # Setup Prompt
    prompt = PromptTemplate(
        template="You are an expert ATS (Applicant Tracking System) software and technical recruiter. "
                 "Analyze the provided Resume text against the Job Description text.\n\n"
                 "Resume:\n{resume_text}\n\n"
                 "Job Description:\n{job_description}\n\n"
                 "{format_instructions}\n",
        input_variables=["resume_text", "job_description"],
        partial_variables={"format_instructions": parser.get_format_instructions()}
    )

    try:
        _input = prompt.format_prompt(resume_text=text, job_description=job_description)
        response = llm.invoke(_input.to_messages())
        analysis_result = parser.parse(response.content)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI processing failed: {str(e)}")
        
    # --- Database Persistence ---
    try:
        db = get_auth_client(current_user.token)
        
        # 1. Insert/Update Resume (We skip physical file upload for now to keep it fast)
        resume_data = {
            "user_id": current_user.id,
            "title": f"Resume for {analysis_result.job_title}",
            "file_path": f"local/{resume.filename}",
            "file_type": "pdf" if resume.filename.endswith('.pdf') else "docx",
            "extracted_text": text
        }
        res_res = db.table("resumes").insert(resume_data).execute()
        resume_id = res_res.data[0]['id']
        
        # 2. Insert Job Description
        job_data = {
            "user_id": current_user.id,
            "company_name": analysis_result.company_name,
            "job_title": analysis_result.job_title,
            "raw_text": job_description
        }
        job_res = db.table("job_descriptions").insert(job_data).execute()
        job_id = job_res.data[0]['id']
        
        # 3. Insert ATS Analysis
        ats_data = {
            "user_id": current_user.id,
            "resume_id": resume_id,
            "job_id": job_id,
            "overall_score": analysis_result.ats_score,
            "matching_skills": analysis_result.matching_skills,
            "missing_skills": analysis_result.missing_skills,
            "ai_summary": analysis_result.summary
        }
        db.table("ats_analyses").insert(ats_data).execute()
    except Exception as e:
        print(f"Database insertion failed: {e}")
        # We don't fail the request if DB insertion fails, to ensure the user still gets the result.
    
    return {
        "status": "success",
        "ats_score": analysis_result.ats_score,
        "matching_skills": analysis_result.matching_skills,
        "missing_skills": analysis_result.missing_skills,
        "summary": analysis_result.summary
    }

class OptimizeResponse(BaseModel):
    optimized_resume: str = Field(description="The heavily tailored and optimized resume content in markdown format.")

@router.post("/optimize")
async def optimize_resume(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
    current_user: UserData = Depends(get_current_user)
):
    if not resume.filename.endswith(('.pdf', '.docx')):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported.")
    
    file_bytes = await resume.read()
    text = ""
    
    if resume.filename.endswith('.pdf'):
        try:
            text = extract_text_from_pdf(file_bytes)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))
    elif resume.filename.endswith('.docx'):
        try:
            text = extract_text_from_docx(file_bytes)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))

    try:
        llm = ChatGoogleGenerativeAI(
            model="gemini-2.5-pro",
            temperature=0.7,
            google_api_key=settings.GEMINI_API_KEY
        )
        parser = PydanticOutputParser(pydantic_object=OptimizeResponse)

        prompt = PromptTemplate(
            template="""You are an elite executive resume writer and ATS optimization expert.
You have been given a candidate's original resume and a target job description.

Your task is to REWRITE the resume so it perfectly aligns with the job description to pass ATS systems and impress recruiters.

Original Resume:
{resume_text}

Target Job Description:
{job_description}

Rules:
1. Do NOT invent fake experience or lie. Use the candidate's actual experience but frame it using the exact terminology and keywords from the job description.
2. Emphasize metrics and achievements (e.g., "Increased sales by 20%"). If none exist, suggest placeholders like [Metric].
3. Ensure the summary/objective clearly targets the role.
4. Output the complete optimized resume in highly readable Markdown format. Include sections for Summary, Skills, Experience, and Education.

{format_instructions}
""",
            input_variables=["resume_text", "job_description"],
            partial_variables={"format_instructions": parser.get_format_instructions()}
        )
        
        chain = prompt | llm | parser
        result = chain.invoke({
            "resume_text": text,
            "job_description": job_description
        })

        return {"status": "success", "optimized_resume": result.optimized_resume}

    except Exception as e:
        error_msg = str(e)
        if "RESOURCE_EXHAUSTED" in error_msg or "429" in error_msg or "quota" in error_msg.lower():
            error_msg = "Gemini API Rate Limit Exceeded. Please wait a minute before trying again or upgrade your API key tier."
        print(f"Resume Optimization Error: {error_msg}")
        raise HTTPException(status_code=500, detail=error_msg)
